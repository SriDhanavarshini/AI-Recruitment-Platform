from __future__ import annotations

from datetime import datetime, timezone
from io import BytesIO
from types import SimpleNamespace
from uuid import UUID

import pytest
from docx import Document
from fastapi.testclient import TestClient

from app.api.deps import get_current_user
from app.api.routers import ats
from app.main import app
from app.schemas.ats_schemas import ResumeInformation
from app.services import ats_pipeline_service, resume_extraction_service
from app.services.ai_service import ExternalAIProvider
from app.services.ats_pipeline_service import ATSPipelineService
from app.services.resume_extraction_service import ResumeExtractionError, extract_resume_text

client = TestClient(app)
CANDIDATE_ID = "11111111-1111-4111-8111-111111111111"
OTHER_CANDIDATE_ID = "22222222-2222-4222-8222-222222222222"
HR_ID = "33333333-3333-4333-8333-333333333333"
OTHER_HR_ID = "44444444-4444-4444-8444-444444444444"
COMPANY_ID = "55555555-5555-4555-8555-555555555555"
OTHER_COMPANY_ID = "66666666-6666-4666-8666-666666666666"
JOB_ID = "77777777-7777-4777-8777-777777777777"
APPLICATION_ID = "88888888-8888-4888-8888-888888888888"
RESUME_ID = "99999999-9999-4999-8999-999999999999"
RESULT_ID = "aaaaaaaa-aaaa-4aaa-8aaa-aaaaaaaaaaaa"
RESUME_TEXT = "Python SQL BSc Computer Science Built Python APIs Inventory platform project."


def make_pdf(text: str) -> bytes:
    stream = f"BT /F1 12 Tf 10 50 Td ({text}) Tj ET".encode("ascii")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 300 100] /Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>",
        b"<< /Length " + str(len(stream)).encode("ascii") + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    pdf = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for index, body in enumerate(objects, start=1):
        offsets.append(len(pdf))
        pdf.extend(f"{index} 0 obj\n".encode("ascii") + body + b"\nendobj\n")
    xref_offset = len(pdf)
    pdf.extend(f"xref\n0 {len(offsets)}\n".encode("ascii"))
    pdf.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        pdf.extend(f"{offset:010d} 00000 n \n".encode("ascii"))
    pdf.extend(
        f"trailer\n<< /Size {len(offsets)} /Root 1 0 R >>\nstartxref\n{xref_offset}\n%%EOF".encode("ascii")
    )
    return bytes(pdf)


def make_docx(text: str) -> bytes:
    document = Document()
    document.add_paragraph(text)
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def test_pdf_text_extraction():
    assert "Python backend engineer" in extract_resume_text("resume.pdf", make_pdf("Python backend engineer"))


def test_docx_text_extraction_includes_paragraphs_and_tables():
    document = Document()
    document.add_paragraph("Python engineer")
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Built SQL services"
    output = BytesIO()
    document.save(output)

    text = extract_resume_text("resume.docx", output.getvalue())

    assert "Python engineer" in text
    assert "Built SQL services" in text


@pytest.mark.parametrize(
    ("file_name", "contents", "message"),
    [
        ("resume.txt", b"text", "Unsupported"),
        ("resume.pdf", b"", "empty"),
        ("resume.docx", b"not a docx", "corrupted"),
    ],
)
def test_invalid_resume_files_are_rejected(file_name, contents, message):
    with pytest.raises(ResumeExtractionError, match=message):
        extract_resume_text(file_name, contents)


def test_pdf_with_no_extractable_text_is_rejected():
    with pytest.raises(ResumeExtractionError, match="No extractable text"):
        extract_resume_text("resume.pdf", make_pdf(""))


def test_ai_response_requires_all_structured_fields():
    with pytest.raises(ValueError):
        ResumeInformation.model_validate({"skills": ["Python"], "education": []})


def test_ai_response_rejects_unexpected_fields():
    with pytest.raises(ValueError):
        ResumeInformation.model_validate({
            "skills": [], "education": [], "experience": "", "projects": [], "score": 100,
        })


def test_gemini_extraction_requests_structured_json_and_validates_response():
    calls = []

    class MockModels:
        def generate_content(self, **kwargs):
            calls.append(kwargs)
            return SimpleNamespace(text=(
                '{"skills":["Python"],"education":[],"experience":"Built Python APIs","projects":[]}'
            ))

    provider = ExternalAIProvider()
    provider.api_key = "test-key"
    provider._client = SimpleNamespace(models=MockModels())

    result = provider.extract_resume_information("Built Python APIs")

    assert result.skills == ["Python"]
    assert calls[0]["model"] == provider.model
    assert calls[0]["config"].response_mime_type == "application/json"
    assert calls[0]["config"].response_schema is ResumeInformation


class FakeQuery:
    def __init__(self, database, table_name):
        self.database = database
        self.table_name = table_name
        self.action = "select"
        self.payload = None
        self.filters = {}
        self.limit_value = None
        self.single_result = False

    def select(self, _columns):
        return self

    def insert(self, payload):
        self.action = "insert"
        self.payload = dict(payload)
        return self

    def eq(self, column, value):
        self.filters[column] = str(value)
        return self

    def maybe_single(self):
        self.single_result = True
        return self

    def single(self):
        self.single_result = True
        return self

    def limit(self, value):
        self.limit_value = value
        return self

    def execute(self):
        if self.database.fail_table == self.table_name:
            raise RuntimeError("mock database failure")
        rows = self.database.tables[self.table_name]
        if self.action == "insert":
            self.database.writes.append((self.table_name, self.payload))
            if self.table_name == "ats_results" and any(
                row["application_id"] == self.payload["application_id"] for row in rows
            ):
                error = RuntimeError("duplicate key value violates unique constraint")
                error.code = "23505"
                raise error
            row = {
                **self.payload,
                "id": RESULT_ID,
                "created_at": "2026-10-02T12:00:00+00:00",
            }
            rows.append(row)
            return SimpleNamespace(data=row)

        matched = [
            row for row in rows
            if all(str(row.get(column)) == value for column, value in self.filters.items())
        ]
        if self.limit_value is not None:
            matched = matched[: self.limit_value]
        if self.single_result:
            return SimpleNamespace(data=matched[0] if matched else None)
        return SimpleNamespace(data=matched)


class FakeStorageBucket:
    def __init__(self, database):
        self.database = database

    def download(self, object_path):
        if self.database.fail_storage:
            raise RuntimeError("mock storage failure")
        self.database.downloaded_paths.append(object_path)
        return self.database.resume_bytes


class FakeStorage:
    def __init__(self, database):
        self.database = database

    def from_(self, _bucket):
        return FakeStorageBucket(self.database)


class FakeDatabase:
    def __init__(self):
        self.tables = {
            "applications": [{
                "id": APPLICATION_ID,
                "candidate_id": CANDIDATE_ID,
                "job_id": JOB_ID,
                "resume_id": RESUME_ID,
                "status": "APPLIED",
                "applied_at": "2026-10-01T12:00:00+00:00",
            }],
            "resumes": [{
                "id": RESUME_ID,
                "candidate_id": CANDIDATE_ID,
                "file_name": "resume.pdf",
                "file_url": (
                    "https://storage.example/storage/v1/object/resumes/"
                    f"{CANDIDATE_ID}/resume.pdf"
                ),
            }],
            "jobs": [{
                "id": JOB_ID,
                "company_id": COMPANY_ID,
                "title": "Backend Engineer",
                "description": "Build Python and SQL services",
                "required_skills": ["Python", "SQL", "Kubernetes"],
                "required_experience": "3 years backend experience",
            }],
            "companies": [{"id": COMPANY_ID, "created_by": HR_ID}],
            "ats_results": [],
        }
        self.resume_bytes = make_pdf(RESUME_TEXT)
        self.storage = FakeStorage(self)
        self.writes = []
        self.downloaded_paths = []
        self.fail_table = None
        self.fail_storage = False

    def table(self, table_name):
        return FakeQuery(self, table_name)


class FakeInformationExtractor:
    def __call__(self, _resume_text):
        return {
            "skills": ["Python", "SQL", "invented qualification"],
            "education": ["BSc Computer Science"],
            "experience": "Built Python APIs",
            "projects": ["Inventory platform project", "fabricated project"],
        }


def fixed_score(_resume_text, _job_description):
    return {
        "score": 81.4,
        "semantic": 90.0,
        "keyword": 68.5,
        "matched_keywords": ["python", "sql"],
        "missing_keywords": ["kubernetes"],
    }


def make_service(database=None, info_extractor=None):
    database = database or FakeDatabase()
    service = ATSPipelineService(
        database=database,
        information_extractor=info_extractor or FakeInformationExtractor(),
        scorer=fixed_score,
    )
    return service, database


def candidate():
    return {"id": CANDIDATE_ID, "role": "CANDIDATE"}


def hr(user_id=HR_ID):
    return {"id": user_id, "role": "HR"}


def test_pipeline_validates_relationships_extracts_scores_and_stores_result(monkeypatch):
    service, database = make_service()
    result = service.process_application(UUID(APPLICATION_ID), candidate())

    assert result["score"] == 81.4
    assert result["semantic"] == 90.0
    assert result["keyword"] == 68.5
    assert result["matching_skills"] == ["Python", "SQL"]
    assert result["missing_skills"] == ["Kubernetes"]
    assert result["extracted_skills"] == ["Python", "SQL"]
    assert result["extracted_projects"] == ["Inventory platform project"]
    assert "invented qualification" not in result["explanation"]
    assert "fabricated project" not in result["explanation"]
    assert "Relevant experience evidence" in result["explanation"]
    assert database.downloaded_paths == [f"{CANDIDATE_ID}/resume.pdf"]
    assert len(database.writes) == 1
    table_name, payload = database.writes[0]
    assert table_name == "ats_results"
    assert set(payload) == {
        "application_id", "score", "matching_skills", "missing_skills", "extracted_skills",
        "extracted_education", "extracted_experience", "extracted_projects", "explanation",
    }


def test_candidate_cannot_process_another_candidates_application():
    service, database = make_service()
    database.tables["applications"][0]["candidate_id"] = OTHER_CANDIDATE_ID

    from app.services.ats_pipeline_service import ATSNotFoundError
    with pytest.raises(ATSNotFoundError):
        service.process_application(UUID(APPLICATION_ID), candidate())
    assert not database.downloaded_paths
    assert not database.writes


def test_resume_must_belong_to_application_candidate():
    service, database = make_service()
    database.tables["resumes"][0]["candidate_id"] = OTHER_CANDIDATE_ID

    from app.services.ats_pipeline_service import ATSNotFoundError
    with pytest.raises(ATSNotFoundError):
        service.process_application(UUID(APPLICATION_ID), candidate())
    assert not database.downloaded_paths


def test_application_job_must_exist():
    service, database = make_service()
    database.tables["applications"][0]["job_id"] = OTHER_COMPANY_ID

    from app.services.ats_pipeline_service import ATSNotFoundError
    with pytest.raises(ATSNotFoundError, match="job"):
        service.process_application(UUID(APPLICATION_ID), candidate())
    assert not database.downloaded_paths
    assert not database.writes


def test_hr_can_process_only_company_owned_applications():
    service, database = make_service()
    result = service.process_application(UUID(APPLICATION_ID), hr())
    assert result["application_id"] == APPLICATION_ID

    other_service, other_database = make_service()
    other_database.tables["companies"][0]["created_by"] = OTHER_HR_ID
    from app.services.ats_pipeline_service import ATSNotFoundError
    with pytest.raises(ATSNotFoundError):
        other_service.process_application(UUID(APPLICATION_ID), hr())
    assert not other_database.downloaded_paths


def test_duplicate_ats_result_is_rejected():
    service, database = make_service()
    database.tables["ats_results"].append({"application_id": APPLICATION_ID})

    from app.services.ats_pipeline_service import ATSConflictError
    with pytest.raises(ATSConflictError):
        service.process_application(UUID(APPLICATION_ID), candidate())
    assert not database.downloaded_paths
    assert not database.writes


def test_unique_constraint_race_maps_to_duplicate_conflict():
    service, database = make_service()

    class RaceQuery(FakeQuery):
        def execute(self):
            if self.table_name == "ats_results" and self.action == "insert":
                error = RuntimeError("duplicate key value violates unique constraint")
                error.code = "23505"
                raise error
            return super().execute()

    database.table = lambda table_name: RaceQuery(database, table_name)
    from app.services.ats_pipeline_service import ATSConflictError
    with pytest.raises(ATSConflictError):
        service.process_application(UUID(APPLICATION_ID), candidate())


def test_invalid_ai_result_is_rejected_and_not_stored():
    service, database = make_service(info_extractor=lambda _text: {"skills": ["Python"]})

    from app.services.ats_pipeline_service import ATSProcessingError
    with pytest.raises(ATSProcessingError, match="invalid resume information"):
        service.process_application(UUID(APPLICATION_ID), candidate())
    assert not database.writes


def test_ai_failure_is_reported_without_database_write():
    def fail_ai(_text):
        raise RuntimeError("mocked AI failure")

    service, database = make_service(info_extractor=fail_ai)
    from app.services.ats_pipeline_service import ATSProcessingError
    with pytest.raises(ATSProcessingError, match="information extraction failed"):
        service.process_application(UUID(APPLICATION_ID), candidate())
    assert not database.writes


def test_supabase_and_storage_failures_are_mapped():
    service, database = make_service()
    database.fail_table = "applications"
    from app.services.ats_pipeline_service import ATSProcessingError
    with pytest.raises(ATSProcessingError, match="validate application relationships"):
        service.process_application(UUID(APPLICATION_ID), candidate())

    service, database = make_service()
    database.fail_storage = True
    with pytest.raises(ATSProcessingError, match="retrieve resume"):
        service.process_application(UUID(APPLICATION_ID), candidate())
    assert not database.writes


def test_resume_storage_path_must_match_application_candidate():
    service, database = make_service()
    database.tables["resumes"][0]["file_url"] = (
        "https://storage.example/storage/v1/object/resumes/"
        f"{OTHER_CANDIDATE_ID}/resume.pdf"
    )

    from app.services.ats_pipeline_service import ATSProcessingError
    with pytest.raises(ATSProcessingError, match="does not match"):
        service.process_application(UUID(APPLICATION_ID), candidate())
    assert not database.downloaded_paths
    assert not database.writes


def test_api_authentication_and_role_access(monkeypatch):
    database = FakeDatabase()
    monkeypatch.setattr(ats_pipeline_service, "supabase", database)
    monkeypatch.setattr(ats_pipeline_service, "extract_resume_information", FakeInformationExtractor())
    monkeypatch.setattr(ats_pipeline_service, "ats_score", fixed_score)
    route = f"/ats/applications/{APPLICATION_ID}/process"

    assert client.post(route).status_code == 401

    app.dependency_overrides[get_current_user] = lambda: {"id": HR_ID, "role": "HR"}
    try:
        response = client.post(route)
    finally:
        app.dependency_overrides.pop(get_current_user, None)
    assert response.status_code == 200

    app.dependency_overrides[get_current_user] = lambda: {"id": OTHER_CANDIDATE_ID, "role": "CANDIDATE"}
    try:
        forbidden = client.post(route)
    finally:
        app.dependency_overrides.pop(get_current_user, None)
    assert forbidden.status_code == 404


def test_api_maps_ai_and_supabase_failures_to_safe_responses(monkeypatch):
    database = FakeDatabase()
    monkeypatch.setattr(ats_pipeline_service, "supabase", database)
    monkeypatch.setattr(
        ats_pipeline_service,
        "extract_resume_information",
        lambda _text: (_ for _ in ()).throw(RuntimeError("private provider detail")),
    )
    app.dependency_overrides[get_current_user] = lambda: {"id": CANDIDATE_ID, "role": "CANDIDATE"}
    try:
        response = client.post(f"/ats/applications/{APPLICATION_ID}/process")
    finally:
        app.dependency_overrides.pop(get_current_user, None)
    assert response.status_code == 503
    assert "private provider detail" not in response.text


def test_ats_result_can_be_read_only_by_authorized_owner(monkeypatch):
    database = FakeDatabase()
    result_row = {
        "id": RESULT_ID,
        "application_id": APPLICATION_ID,
        "score": 81.4,
        "matching_skills": ["Python"],
        "missing_skills": ["Kubernetes"],
        "extracted_skills": ["Python"],
        "extracted_education": [],
        "extracted_experience": "Built Python APIs",
        "extracted_projects": [],
        "explanation": "Score details",
        "created_at": "2026-10-02T12:00:00+00:00",
    }
    database.tables["ats_results"].append(result_row)
    monkeypatch.setattr(ats_pipeline_service, "supabase", database)
    app.dependency_overrides[get_current_user] = lambda: {"id": CANDIDATE_ID, "role": "CANDIDATE"}
    try:
        response = client.get(f"/ats/applications/{APPLICATION_ID}")
    finally:
        app.dependency_overrides.pop(get_current_user, None)
    assert response.status_code == 200
    assert response.json()["id"] == RESULT_ID


def test_api_router_is_the_only_http_entry_for_pipeline(monkeypatch):
    database = FakeDatabase()
    monkeypatch.setattr(ats_pipeline_service, "supabase", database)
    monkeypatch.setattr(ats, "ATSPipelineService", lambda: ATSPipelineService(
        database=database,
        information_extractor=FakeInformationExtractor(),
        scorer=fixed_score,
    ))
    app.dependency_overrides[get_current_user] = lambda: {"id": CANDIDATE_ID, "role": "CANDIDATE"}
    try:
        response = client.post(f"/ats/applications/{APPLICATION_ID}/process")
    finally:
        app.dependency_overrides.pop(get_current_user, None)
    assert response.status_code == 200
    assert response.json()["score"] == 81.4
