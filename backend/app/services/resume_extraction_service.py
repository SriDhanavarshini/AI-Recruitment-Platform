from __future__ import annotations

from io import BytesIO
from pathlib import Path


class ResumeExtractionError(ValueError):
    pass


def extract_resume_text(file_name: str, file_bytes: bytes) -> str:
    if not file_bytes:
        raise ResumeExtractionError("Resume file is empty")

    extension = Path(file_name).suffix.lower()
    if extension not in {".pdf", ".docx"}:
        raise ResumeExtractionError("Unsupported resume format; upload a PDF or DOCX file")

    try:
        if extension == ".pdf":
            text = _extract_pdf(file_bytes)
        else:
            text = _extract_docx(file_bytes)
    except ResumeExtractionError:
        raise
    except Exception:
        raise ResumeExtractionError("Resume file is corrupted or could not be read") from None

    text = "\n".join(line.strip() for line in text.splitlines() if line.strip()).strip()
    if not text:
        raise ResumeExtractionError("No extractable text found in resume")
    return text


def _extract_pdf(file_bytes: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(BytesIO(file_bytes))
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def _extract_docx(file_bytes: bytes) -> str:
    from docx import Document

    document = Document(BytesIO(file_bytes))
    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    table_text = [cell.text for table in document.tables for row in table.rows for cell in row.cells]
    return "\n".join(paragraphs + table_text)